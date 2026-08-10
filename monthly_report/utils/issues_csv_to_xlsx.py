import sys
import csv
import os
import re
import argparse
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Updated to match the exact wording requested for both hover comments and the Penjelasan Header sheet
COMMENTS = {
    "project_name": "project_name\nNama project di SonarQube, sesuai yang di configure di UI",
    "project_key": "project_key\nKey unique yang untuk mengidentifikasi project, digunakan di API call dan URL",
    "issue_key": "issue_key\nIdentifier unique secara global untuk setiap issue atau security hotspot",
    "type": "type\nKategori temuan: reliability_issue (bug), maintainability_issue (code smell), vulnerability (kerentanan), atau SECURITY_HOTSPOT (area sensitif yang perlu review keamanan manual)",
    "severity": "severity\nLevel dampak dari temuan: CRITICAL, HIGH, MEDIUM, LOW, atau INFO",
    "rule": "rule\nRule ID di SonarQube yang memicu temuan ini (contoh: java:S2095)",
    "component": "component\nPath lengkap ke source file yang terdapat issue, diawali dengan project key",
    "line": "line\nBaris pada file dimana issue ditemukan; N/A jika temuan berlaku untuk keseluruhan file",
    "message": "message\nDeskripsi singkat dari rule engine SonarQube yang menjelaskan apa yang salah pada temuan tersebut",
    "committer": "committer\nGit author (email atau login) yang melakukan commit kode, berdasarkan data SCM blame",
    "assigner": "assigner\nDisplay name user yang melakukan assign issue ke orang lain, diambil dari changelog. Kosong jika issue di-assign otomatis oleh sistem (bukan oleh user secara manual).",
    "assignee": "assignee\nDisplay name user yang saat ini bertanggung jawab menyelesaikan issue. Kosong jika terjadi email mismatch antara Git commit dan akun SonarQube sehingga auto-assign gagal. Meski kosong, ini tidak mengganggu penyelesaian issue — developer (committer) bisa membuka temuan langsung melalui SonarQube MR comment di GitLab.",
    "resolution": "resolution\nCara issue ditutup: Fixed, False Positive, Accepted, Safe, atau Acknowledged (dua terakhir khusus hotspot)",
    "justification/comment": "justification/comment\nSemua komentar pada issue secara berurutan, format [N] Author (YYYY-MM-DD HH:MM): teks, dipisahkan dengan simbol pipe |",
    "created_at": "created_at\nTimestamp ISO 8601 saat SonarQube pertama kali mendeteksi issue",
    "resolution_date": "resolution_date\nTimestamp ISO 8601 saat issue terakhir diperbarui, yang menandai waktu issue ditutup",
    "days_to_resolve": "days_to_resolve\nLama penyelesaian issue dalam hari, dihitung dari tanggal commit Git pada kode yang dideteksi sebagai issue (bukan dari created_at). Nilai 0 berarti diselesaikan di hari yang sama.",
    "age_days": "age_days\nUmur issue (dalam hari) sejak commit Git yang memicu temuan hingga tanggal capture. Semakin besar nilainya, semakin lama issue belum diselesaikan.",
    "effort": "effort\nEstimasi waktu yang dibutuhkan untuk memperbaiki issue, dihitung oleh SonarQube berdasarkan rule. Format dalam menit (contoh: 30min). N/A untuk SECURITY_HOTSPOT."
}

def extract_date_from_filename(filename):
    """Extracts a date like YYYY-MM-DD or YYYY_MM_DD from the file path and formats it with underscores."""
    base_name = os.path.basename(filename)
    match = re.search(r'(\d{4}[-_]\d{2}[-_]\d{2})', base_name)
    if match:
        return match.group(1).replace('-', '_')
    return "unknown_date"

def calc_size(text):
    width = 220
    chars_per_line = width // 7
    total_lines = 0
    for line in text.split("\n"):
        total_lines += max(1, -(-len(line) // chars_per_line))
    height = min(max(total_lines * 20 + 20, 60), 500)
    return width, height

def col_width(values, header):
    all_vals = [str(header)] + [str(v) for v in values if v]
    ideal = min(max(len(v) for v in all_vals) * 1.15, 60)
    return max(ideal, 10)

def add_csv_to_sheet(csv_path, ws, table_name):
    """Reads a CSV and formats it into the provided worksheet."""
    if not os.path.isfile(csv_path):
        print(f"Warning: File '{csv_path}' not found. Leaving sheet empty.")
        return

    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if not rows:
        print(f"CSV '{csv_path}' is empty.")
        return

    headers = rows[0]
    data = rows[1:]

    # Write headers
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="2E75B6")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

        # Inject hover comments on headers
        col_name = header.strip()
        if col_name in COMMENTS:
            text = COMMENTS[col_name]
            comment = Comment(text, "SonarQube Audit")
            comment.width, comment.height = calc_size(text)
            cell.comment = comment

    # Write data rows
    for row_idx, row in enumerate(data, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    # Excel Table formatting
    last_col = get_column_letter(len(headers))
    last_row = len(rows)
    table_ref = f"A1:{last_col}{last_row}"
    table = Table(displayName=table_name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    # Auto column widths
    col_data = {col_idx: [] for col_idx in range(1, len(headers) + 1)}
    for row in data:
        for col_idx, value in enumerate(row, start=1):
            col_data[col_idx].append(value)
    
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width(col_data[col_idx], header)

    # Freeze header row
    ws.freeze_panes = "A2"


def generate_glossary_sheet(ws):
    """Generates the Penjelasan Header tab."""
    # Custom headers for the glossary
    ws.append(["Kolom", "Keterangan"])
    
    # Format glossary headers
    for col in range(1, 3):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", start_color="2E75B6")
        cell.alignment = Alignment(horizontal="left", vertical="center")

    # Populate rows
    for row_idx, (col_name, full_comment) in enumerate(COMMENTS.items(), start=2):
        # Extract just the description part (everything after the first newline)
        description = full_comment.split("\n", 1)[-1] if "\n" in full_comment else full_comment
        
        ws.cell(row=row_idx, column=1, value=col_name)
        desc_cell = ws.cell(row=row_idx, column=2, value=description)
        desc_cell.alignment = Alignment(wrap_text=True, vertical="top")
        
    # Set widths for readability
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 120


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Merge Open and Closed CSVs into a single Excel file.")
    parser.add_argument("--open", required=True, help="Path to the open issues CSV")
    parser.add_argument("--closed", required=True, help="Path to the closed issues CSV")
    parser.add_argument("--output", required=True, help="Path for the output Excel file (e.g., Master_Report.xlsx)")
    args = parser.parse_args()

    # Extract dates dynamically from the filenames
    open_date = extract_date_from_filename(args.open)
    closed_date = extract_date_from_filename(args.closed)

    # Initialize a new workbook
    wb = Workbook()
    
    # 1. Open Issues Sheet
    ws_open = wb.active
    ws_open.title = f"open_issues_{open_date}"
    print(f"Processing Open Issues from: {args.open}")
    add_csv_to_sheet(args.open, ws_open, "OpenIssuesTable")

    # 2. Closed Issues Sheet
    ws_closed = wb.create_sheet(title=f"closed_issues_{closed_date}")
    print(f"Processing Closed Issues from: {args.closed}")
    add_csv_to_sheet(args.closed, ws_closed, "ClosedIssuesTable")

    # 3. Glossary Sheet (Penjelasan Header)
    ws_glossary = wb.create_sheet(title="Penjelasan Header")
    print("Generating Penjelasan Header sheet...")
    generate_glossary_sheet(ws_glossary)

    # Save the final file
    wb.save(args.output)
    print(f"Success! Final report saved to: {args.output}")
